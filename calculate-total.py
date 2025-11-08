#!/usr/bin/env python3
import openpyxl

# Load workbook
wb = openpyxl.load_workbook('posicao-2025-10-28-08-41-55.xlsx', data_only=True)

print("=== CALCULATING TOTAL VALUE PER SHEET ===\n")

grand_total = 0

for sheet_name in wb.sheetnames:
    sheet = wb[sheet_name]
    print(f"--- Sheet: \"{sheet_name}\" ---")
    
    # Get headers
    headers = [str(cell.value) if cell.value else '' for cell in sheet[1]]
    
    # Find value column
    value_col_idx = None
    for idx, header in enumerate(headers):
        if any(keyword in header for keyword in ['Valor Atualizado', 'Valor Atualizado CURVA', 'Valor Atualizado MTM']):
            value_col_idx = idx
            print(f"Using column: {header}")
            break
    
    if value_col_idx is None:
        print("No value column found, skipping...")
        continue
    
    sheet_total = 0
    count = 0
    
    for row_num in range(2, sheet.max_row + 1):
        row = list(sheet[row_num])
        if row_num - 1 < len(list(sheet.iter_rows())):
            cell_value = row[value_col_idx].value
            if cell_value:
                try:
                    value = float(cell_value) if isinstance(cell_value, (int, float)) else float(str(cell_value).replace(',', '.'))
                    if value > 0:
                        sheet_total += value
                        count += 1
                        print(f"  Row {row_num - 1}: R$ {value:.2f}")
                except:
                    pass
    
    print(f"Sheet Total: R$ {sheet_total:.2f} ({count} assets)")
    print()
    grand_total += sheet_total

print("=" * 50)
print(f"GRAND TOTAL: R$ {grand_total:.2f}")
print("=" * 50)

wb.close()
