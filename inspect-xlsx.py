#!/usr/bin/env python3
import sys

try:
    import openpyxl
except ImportError:
    print("Installing openpyxl...")
    import subprocess
    subprocess.check_call([sys.executable, "-m", "pip", "install", "openpyxl"])
    import openpyxl

# Load workbook
wb = openpyxl.load_workbook('posicao-2025-10-28-08-41-55.xlsx', data_only=True)

print("=== SHEET NAMES ===")
for i, name in enumerate(wb.sheetnames, 1):
    print(f"{i}. \"{name}\"")

print("\n=== SHEET DETAILS ===")
for sheet_name in wb.sheetnames:
    sheet = wb[sheet_name]
    print(f"\n--- Sheet: \"{sheet_name}\" ---")
    print(f"Max row: {sheet.max_row}, Max column: {sheet.max_column}")
    
    # Get headers (first row)
    headers = []
    for cell in sheet[1]:
        headers.append(str(cell.value) if cell.value else '')
    print(f"Headers: {headers}")
    
    # Get first 3 data rows
    print("Sample data (first 3 rows):")
    for row_num in range(2, min(5, sheet.max_row + 1)):
        row_data = []
        for cell in sheet[row_num]:
            row_data.append(str(cell.value) if cell.value else '')
        print(f"  Row {row_num - 1}: {row_data}")
    
    # Count non-empty rows
    non_empty = 0
    for row in sheet.iter_rows(min_row=2):
        if any(cell.value for cell in row):
            non_empty += 1
    print(f"Non-empty data rows: {non_empty}")

wb.close()
